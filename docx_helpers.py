import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def apply_page_border(section):
    sectPr = section._sectPr
    xml_str = f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page"><w:top w:val="single" w:sz="16" w:space="24" w:color="000000"/><w:left w:val="single" w:sz="16" w:space="24" w:color="000000"/><w:bottom w:val="single" w:sz="16" w:space="24" w:color="000000"/><w:right w:val="single" w:sz="16" w:space="24" w:color="000000"/></w:pgBorders>'
    pgBorders = parse_xml(xml_str)
    sectPr.append(pgBorders)

def add_page_number_to_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    # Add page field in Word XML
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    p._p.append(fldSimple)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_code_block(doc, filename, code_text):
    # Header label for the file
    p_label = doc.add_paragraph()
    p_label.paragraph_format.space_before = Pt(8)
    p_label.paragraph_format.space_after = Pt(4)
    p_label.paragraph_format.keep_with_next = True
    r_lbl = p_label.add_run(filename)
    r_lbl.bold = True
    r_lbl.font.size = Pt(11)
    r_lbl.font.name = "Consolas"
    r_lbl.font.color.rgb = RGBColor(15, 23, 42)

    # Table wrapper for the code
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/><w:left w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/><w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_diagram_box(doc, title, diagram_str):
    p_label = doc.add_paragraph()
    p_label.paragraph_format.space_before = Pt(10)
    p_label.paragraph_format.space_after = Pt(4)
    p_label.paragraph_format.keep_with_next = True
    r = p_label.add_run(f"Figure: {title}")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(15, 23, 42)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="12" w:space="0" w:color="94A3B8"/><w:left w:val="single" w:sz="12" w:space="0" w:color="94A3B8"/><w:bottom w:val="single" w:sz="12" w:space="0" w:color="94A3B8"/><w:right w:val="single" w:sz="12" w:space="0" w:color="94A3B8"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(diagram_str)
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(15, 23, 42)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

print("Helper functions defined.")
